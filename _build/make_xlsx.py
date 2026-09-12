# -*- coding: utf-8 -*-
"""통합 엑셀 자료 생성 (openpyxl). 글꼴 크기는 모두 보통 셀 서식으로만 지정합니다."""
import os
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.hyperlink import Hyperlink
from chemdata import ELEMENTS, MOLECULES, CATIONS, ANIONS, COMPOUNDS, RULES, plain, is_polyatomic

HERE = os.path.dirname(os.path.abspath(__file__))
OUT_DIR = os.path.join(HERE, "out")
OUT = os.path.join(OUT_DIR, "workbook.xlsx")

FONT = "맑은 고딕"
NAVY, BLUE, GREEN, CORAL, VIOLET, AMBER, MUTED, INK = "1F3A5F", "0F6FB8", "1F7A68", "C4404B", "7A3FB0", "B7791F", "6B7684", "1A2433"
YELLOW, SOFT, WHITE, LINE, VIOLET_BG = "FFF3C4", "F3F6FA", "FFFFFF", "D8E0EA", "F3EBFB"

def font(sz=11, bold=False, color=INK):
    return Font(name=FONT, size=sz, bold=bold, color=color)
def fill(rgb): return PatternFill("solid", fgColor=rgb)
thin = Side(style="thin", color=LINE)
BORDER = Border(left=thin, right=thin, top=thin, bottom=thin)
CENTER = Alignment(horizontal="center", vertical="center", wrap_text=True)
LEFT = Alignment(horizontal="left", vertical="center", wrap_text=True, indent=1)

wb = Workbook()
wb.remove(wb.active)

def new_sheet(title, tab, widths, page_title, subtitle):
    ws = wb.create_sheet(title)
    ws.sheet_properties.tabColor = tab
    ws.sheet_view.showGridLines = False
    ws.column_dimensions["A"].width = 2.2
    for i, w in enumerate(widths, start=2):
        ws.column_dimensions[get_column_letter(i)].width = w
    last = get_column_letter(len(widths) + 1)
    for r, h in ((1, 8), (2, 30), (3, 18), (4, 6)):
        ws.row_dimensions[r].height = h
    ws.merge_cells(f"B2:{last}2"); ws.merge_cells(f"B3:{last}3")
    ws["B2"] = page_title; ws["B2"].font = font(17, True, NAVY); ws["B2"].alignment = Alignment(vertical="center")
    ws["B3"] = subtitle;   ws["B3"].font = font(10, False, MUTED); ws["B3"].alignment = Alignment(vertical="center")
    return ws

def header(ws, row, labels, height=30):
    ws.row_dimensions[row].height = height
    for i, label in enumerate(labels, start=2):
        c = ws.cell(row=row, column=i, value=label)
        c.font = font(10, True, WHITE); c.fill = fill(NAVY); c.alignment = CENTER; c.border = BORDER

def write_row(ws, row, values, highlight=False, height=20, fonts=None, aligns=None, formats=None, bg=None):
    ws.row_dimensions[row].height = height
    bg = bg or (YELLOW if highlight else (SOFT if row % 2 == 0 else WHITE))
    for i, v in enumerate(values, start=2):
        c = ws.cell(row=row, column=i, value=v)
        c.font = (fonts or {}).get(i - 2, font(11))
        c.alignment = (aligns or {}).get(i - 2, CENTER)
        c.fill = fill(bg); c.border = BORDER
        if formats and (i - 2) in formats:
            c.number_format = formats[i - 2]

def notes(ws, start_row, lines, last_col):
    r = start_row
    for line in lines:
        ws.merge_cells(start_row=r, start_column=2, end_row=r, end_column=last_col)
        c = ws.cell(row=r, column=2, value=line); c.font = font(9, False, MUTED); c.alignment = Alignment(vertical="center", wrap_text=True)
        ws.row_dimensions[r].height = 18
        r += 1
    return r

def pages_font(pages, sz=9):
    return font(sz, False, "8A6D1F" if pages != "–" else MUTED)

# ------------------------------------------------------------------ 원소 기호
core_el = sum(1 for e in ELEMENTS if e["core"])
ws = new_sheet("원소 기호", "2E5A88", [11, 22, 12, 11, 12, 20], "원소 기호",
               f"원소 {len(ELEMENTS)}종 중 학습 핵심 {core_el}종 (연노란색 · 학습지 4쪽에 수록) · 원자 번호 순")
header(ws, 5, ["원자 번호", "원소 이름", "원소 기호", "학습 범위", "학습지 수록", "교과서 확인"])
r = 6
for e in ELEMENTS:
    label = e["name"] + (f"({'·'.join(e['alias'])})" if e["alias"] else "")
    write_row(ws, r, [e["z"], label, e["symbol"], "핵심" if e["core"] else "확장", "○" if e["core"] else "–", e["pages"]], highlight=e["core"],
              fonts={0: font(10, False, MUTED), 1: font(11), 2: font(12, True, NAVY), 3: font(10, True, BLUE if e["core"] else MUTED),
                     4: font(10, False, INK if e["core"] else MUTED), 5: pages_font(e["pages"])})
    r += 1
notes(ws, r + 1, [
    f"※ 연노란색 = 학습 핵심 {core_el}종. 학습 사이트의 '교과서·학습 핵심만' 범위이자 4쪽 학습지에 실린 원소입니다.",
    "※ 원자 번호 1~20번 외에 망가니즈(25)·철(26)·구리(29)·아연(30)·브로민(35)·은(47)·아이오딘(53)·금(79)·수은(80)·납(82)을 핵심에 넣었습니다.",
    "※ 교과서 확인 쪽수는 Ⅳ. 물질의 구성 단원(134~165쪽) 기준입니다. 140쪽 = 그림 Ⅳ-2 원소 카드, 145쪽 = 그림 Ⅳ-8 주기율표. '–'는 이 단원에 기호로 나오지 않는 원소입니다.",
    "※ 원소 기호는 첫 글자를 대문자로, 둘째 글자를 소문자로 씁니다. 괄호 안은 IUPAC 이름입니다. 예) 나트륨(소듐), 칼륨(포타슘)",
], 7)
ws.freeze_panes = "C6"

# ------------------------------------------------------------------ 분자식
core_mol = sum(1 for m in MOLECULES if m["core"])
ws = new_sheet("분자식", GREEN, [7, 14, 14, 18, 10, 11, 20], "분자식",
               f"분자식 {len(MOLECULES)}종 중 교과서·학습 핵심 {core_mol}종 (연노란색) · 초록 글자 = 원소, 남색 글자 = 화합물")
header(ws, 5, ["번호", "분자식", "입력용 표기", "물질 이름", "구분", "학습 범위", "교과서 확인"])
r = 6
for i, m in enumerate(MOLECULES, start=1):
    color = GREEN if m["kind"] == "원소" else NAVY
    write_row(ws, r, [i, m["formula"], plain(m["formula"]), m["name"], m["kind"], "핵심" if m["core"] else "확장", m["pages"]], highlight=m["core"],
              fonts={0: font(10, False, MUTED), 1: font(12, True, color), 2: font(10, False, MUTED), 3: font(11, False, color), 4: font(10, True, color),
                     5: font(10, True, BLUE if m["core"] else MUTED), 6: pages_font(m["pages"])})
    r += 1
mol_list = lambda ms: " · ".join(f"{m['name']}({m['formula']})" for m in ms)
notes(ws, r + 1, [
    f"※ 연노란색 = 학습 핵심 {core_mol}종(학습 사이트 '교과서·학습 핵심만' 범위). 교과서 확인 쪽수가 적힌 물질은 Ⅳ단원에 화학식이 실제로 나옵니다. 141쪽 = 그림 Ⅳ-3, 153쪽 = 분자의 화학식.",
    "※ 원소(초록): 한 가지 원소로만 이루어진 물질(H₂, O₂, O₃, N₂) / 화합물(남색): 두 가지 이상의 원소로 이루어진 물질. 원자 수는 아래첨자로 씁니다.",
    f"※ 이 단원에 나오지 않지만 수업에서 자주 다루는 물질 — 학습 핵심: {mol_list(m for m in MOLECULES if m['core'] and m['pages'] == '–')}  /  확장: {mol_list(m for m in MOLECULES if not m['core'])}",
    f"※ 4쪽 학습지(PDF)의 분자식 칸에는 핵심·확장 구분 없이 {len(MOLECULES)}종이 모두 실려 있습니다.",
    "※ '입력용 표기'는 학습 사이트에서 답을 칠 때 쓰는 형태입니다. 숫자는 자동으로 아래첨자가 됩니다. 예) H2O → H₂O",
], 8)
ws.freeze_panes = "C6"

# ------------------------------------------------------------------ 양이온 / 음이온
ION_HEAD = ["번호", "이온의 이름", "이온식", "입력용 표기", "이온의 종류", "원소 이름", "원소 기호", "원자 번호\n(= 양성자 수)", "전하", "잃거나 얻은\n전자 수", "이온의 전자 수", "전류를 흘릴 때\n이동하는 극", "학습 범위", "교과서 수록", "교과서 확인", "비고 · 이름 짓는 법"]
ION_W = [6, 18, 11, 11, 12, 11, 9, 12, 7, 11, 12, 14, 9, 10, 18, 44]

def ion_sheet(title, tab, rows, subtitle, text_color):
    ws = new_sheet(title, tab, ION_W, title, subtitle)
    header(ws, 5, ION_HEAD, height=36)
    r = 6
    for i, ion in enumerate(rows, start=1):
        poly = is_polyatomic(ion["formula"])
        vals = [i, ion["name"], ion["formula"], plain(ion["formula"]), "다원자 이온" if poly else "단원자 이온", ion["element_name"], ion["symbol"],
                ion["z"] if ion["z"] else "-", ion["charge"], f"=ABS(J{r})", f'=IF(I{r}="-","-",I{r}-J{r})', f'=IF(J{r}>0,"(−)극","(+)극")',
                "핵심" if ion["core"] else "확장", ion["textbook"], ion["pages"], ion["note"]]
        write_row(ws, r, vals, highlight=ion["core"], height=21, bg=(VIOLET_BG if poly else None),
                  fonts={0: font(10, False, MUTED), 1: font(11, False, text_color), 2: font(12, True, text_color), 3: font(10, False, MUTED),
                         4: font(10, True, VIOLET if poly else INK), 5: font(10), 6: font(10), 7: font(10), 8: font(11, True, text_color),
                         9: font(10), 10: font(10), 11: font(10), 12: font(10, True, BLUE if ion["core"] else MUTED),
                         13: font(10, False, INK if ion["textbook"] == "교과서" else MUTED), 14: pages_font(ion["pages"]), 15: font(9, False, MUTED)},
                  aligns={15: LEFT}, formats={8: "+0;-0;0"})
        r += 1
    notes(ws, r + 1, [
        "※ 연노란색 = 학습 핵심 이온(학습 사이트 '교과서·학습 핵심만' 범위 · 학습지 수록). 154쪽 그림 Ⅳ-11, 155쪽 그림 Ⅳ-12·표 Ⅳ-1, 157쪽 그림 Ⅳ-13에서 확인했습니다.",
        "※ 연보라색 행 = 다원자 이온(여러 원자가 결합한 이온). 원자 번호가 하나로 정해지지 않으므로 '-'로 표시했습니다. 양이온은 파란 글자, 음이온은 빨간 글자입니다.",
        "※ '잃거나 얻은 전자 수', '이온의 전자 수', '이동하는 극'은 수식입니다. 원자 번호나 전하를 바꾸면 자동으로 다시 계산됩니다.",
        "※ '입력용 표기'는 학습 사이트에서 답을 칠 때 쓰는 형태입니다. 맨 끝의 숫자와 부호는 자동으로 위첨자(전하)가 됩니다. 예) Cu2+ → Cu²⁺, SO42- → SO₄²⁻",
    ], 17)
    ws.freeze_panes = "D6"

ion_sheet("양이온", BLUE, CATIONS, f"전자를 잃어 (+)전하를 띤 이온 {len(CATIONS)}종 중 학습 핵심 {sum(1 for c in CATIONS if c['core'])}종 (연노란색) · 연보라색 행 = 다원자 이온", BLUE)
ion_sheet("음이온", CORAL, ANIONS, f"전자를 얻어 (−)전하를 띤 이온 {len(ANIONS)}종 중 학습 핵심 {sum(1 for a in ANIONS if a['core'])}종 (연노란색) · 연보라색 행 = 다원자 이온", CORAL)

# ------------------------------------------------------------------ 이온 이름 규칙
ws = new_sheet("이온 이름 규칙", VIOLET, [7, 52, 52, 26], "이온 이름 규칙", "이온 이름을 붙이는 규칙과 예시 · 잘못된 표기")
header(ws, 5, ["번호", "규칙", "예시", "잘못된 표기"])
r = 6
for i, (rule, ex, wrong) in enumerate(RULES, start=1):
    write_row(ws, r, [i, rule, ex, wrong], height=26, fonts={0: font(10, False, MUTED), 1: font(11), 2: font(11, False, NAVY), 3: font(11, False, CORAL)}, aligns={1: LEFT, 2: LEFT, 3: LEFT})
    r += 1
notes(ws, r + 1, [
    "※ 1~6번 규칙은 교과서 155쪽 그림 Ⅳ-12와 본문, 8번은 157쪽 그림 Ⅳ-13의 내용을 정리한 것입니다.",
    "※ '잘못된 표기'는 학생들이 자주 틀리는 형태입니다. 규칙과 나란히 놓고 비교해 보세요.",
], 5)
ws.freeze_panes = "C6"

# ------------------------------------------------------------------ 이온 결합 화합물
ws = new_sheet("이온 결합 화합물", "D65A64", [6, 20, 11, 11, 10, 12, 11, 10, 9, 9, 13, 13, 40], "이온 결합 화합물",
               f"양이온과 음이온의 전체 전하 합이 0이 되는 가장 간단한 정수비 · {len(COMPOUNDS)}종 (모두 학습 핵심) · 초록 = 개수비 1 : 1, 주황 = 개수비 다름")
header(ws, 5, ["번호", "화합물 이름", "양이온", "양이온 전하", "양이온 수", "음이온", "음이온 전하", "음이온 수", "개수비", "전하 합", "화학식", "입력용 표기", "학습 메모"])
r = 6
for i, c in enumerate(COMPOUNDS, start=1):
    same = c["cation_count"] == c["anion_count"]
    color = GREEN if same else AMBER
    write_row(ws, r, [i, c["name"], c["cation"], c["cation_charge"], c["cation_count"], c["anion"], c["anion_charge"], c["anion_count"], c["ratio"],
                      f"=E{r}*F{r}+H{r}*I{r}", c["formula"], plain(c["formula"]), c["memo"]], highlight=True, height=22,
              fonts={0: font(10, False, MUTED), 1: font(11, True, color), 2: font(11, True, BLUE), 3: font(10), 4: font(10), 5: font(11, True, CORAL), 6: font(10), 7: font(10),
                     8: font(11, True, color), 9: font(10, True, GREEN), 10: font(13, True, color), 11: font(10, False, MUTED), 12: font(9, False, MUTED)},
              aligns={12: LEFT}, formats={3: "+0;-0;0", 6: "+0;-0;0"})
    r += 1
r += 1
ws.merge_cells(start_row=r, start_column=2, end_row=r, end_column=13)
ws.cell(row=r, column=2, value="화학식 만드는 순서").font = font(12, True, NAVY); r += 1
for n, t, d in [("1", "전하 확인", "양이온과 음이온의 전하를 확인합니다."),
                ("2", "최소 정수비", "양전하의 합과 음전하의 합이 같아지는 가장 간단한 정수비를 찾습니다. 전하의 크기가 같으면 1 : 1 입니다."),
                ("3", "아래첨자와 괄호", "이온 수를 아래첨자로 쓰고, 다원자 이온이 2개 이상이면 괄호로 묶습니다. 예) Ca(OH)₂")]:
    ws.merge_cells(start_row=r, start_column=4, end_row=r, end_column=13)
    ws.cell(row=r, column=2, value=n).font = font(10, True, WHITE); ws.cell(row=r, column=2).fill = fill(NAVY); ws.cell(row=r, column=2).alignment = CENTER
    ws.cell(row=r, column=3, value=t).font = font(10, True, NAVY); ws.cell(row=r, column=3).alignment = CENTER
    ws.cell(row=r, column=4, value=d).font = font(10); ws.cell(row=r, column=4).alignment = Alignment(vertical="center", indent=1)
    ws.row_dimensions[r].height = 20; r += 1
notes(ws, r + 1, ["※ '전하 합'은 수식입니다. (양이온 전하 × 양이온 수) + (음이온 전하 × 음이온 수) = 0 이 되어야 합니다.",
                  "※ 초록 글자 = 개수비 1 : 1(전하 크기가 같아 아래첨자 없음), 주황 글자 = 개수비가 달라 아래첨자로 이온 수를 써야 하는 화합물.",
                  "※ 중학교에서는 황산 구리(Ⅱ)를 '황산 구리'로도 표기합니다. 학습 사이트는 두 이름을 모두 정답으로 인정합니다."], 13)
ws.freeze_panes = "C6"

# ------------------------------------------------------------------ 표지 (맨 앞)
cover = wb.create_sheet("표지", 0)
cover.sheet_properties.tabColor = NAVY
cover.sheet_view.showGridLines = False
for col, w in zip("ABCDE", [2.2, 18, 14, 62, 12]):
    cover.column_dimensions[col].width = w
cover.row_dimensions[1].height = 10
cover.merge_cells("B2:E2"); cover["B2"] = "2022 개정 교육과정 중학교 과학 · 꼭 알아둘 화학 기호와 화학식"; cover["B2"].font = font(20, True, NAVY); cover.row_dimensions[2].height = 36
cover.merge_cells("B3:E3"); cover["B3"] = "교과서에 나오는 원소 기호 · 분자식 · 이온식과 이온 결합 화합물의 화학식을 한 파일로 정리했습니다."; cover["B3"].font = font(11, False, MUTED)
cover.merge_cells("B5:E5"); cover["B5"] = "송쌤과학  |  아주중학교  |  YouTube · Instagram @songsamscience"; cover["B5"].font = font(10, True, BLUE)
cover["B7"] = "시트 구성"; cover["B7"].font = font(13, True, NAVY)
header(cover, 8, ["시트", "핵심 / 전체", "내용", "바로가기"], height=24)
index_rows = [
    ("원소 기호", f"{core_el} / {len(ELEMENTS)}종", "원자 번호 · 원소 이름 · 원소 기호 · 학습지 수록 여부"),
    ("분자식", f"{core_mol} / {len(MOLECULES)}종", "물질을 나타내는 분자식 · 원소(초록)와 화합물(남색) 구분"),
    ("양이온", f"{sum(1 for c in CATIONS if c['core'])} / {len(CATIONS)}종", "전자를 잃어 (+)전하를 띤 이온 · 전자 수 자동 계산"),
    ("음이온", f"{sum(1 for a in ANIONS if a['core'])} / {len(ANIONS)}종", "전자를 얻어 (−)전하를 띤 이온 · 다원자 이온은 연보라색 행"),
    ("이온 이름 규칙", f"{len(RULES)}개", "이온 이름을 붙이는 규칙과 예시 · 잘못된 표기 (155쪽 기준)"),
    ("이온 결합 화합물", f"{len(COMPOUNDS)}종", "구성 이온 · 최소 정수비 · 전하 합 · 화학식 (개수비 1 : 1은 초록, 다르면 주황)"),
]
r = 9
for title, count, desc in index_rows:
    write_row(cover, r, [title, count, desc, "이동 →"], height=22, fonts={0: font(11, True, NAVY), 1: font(10), 2: font(10), 3: font(10, True, BLUE)}, aligns={2: LEFT})
    cover.cell(row=r, column=5).hyperlink = Hyperlink(ref=f"E{r}", location=f"'{title}'!A1", display="이동 →", tooltip=f"{title} 시트로 이동")
    r += 1
r += 1
cover.cell(row=r, column=2, value="표 보는 방법").font = font(13, True, NAVY); r += 1
header(cover, r, ["항목", "설명", "", ""], height=22); cover.merge_cells(start_row=r, start_column=3, end_row=r, end_column=5); r += 1
for k, v in [("이온식", "위첨자 · 아래첨자를 그대로 살린 실제 표기입니다. 예) Cu²⁺, SO₄²⁻"),
             ("입력용 표기", "컴퓨터로 입력할 때 쓰는 형태입니다. 학습 사이트는 이 형태로 쳐도 자동으로 첨자를 붙여 줍니다. 예) Cu2+ → Cu²⁺"),
             ("전하", "잃은 전자 수만큼 (+), 얻은 전자 수만큼 (−). 양수는 양이온, 음수는 음이온입니다."),
             ("이온의 전자 수", "원자 번호(양성자 수)에서 전하를 뺀 값으로 계산식이 걸려 있습니다. 다원자 이온은 '-'로 표시합니다."),
             ("이동하는 극", "전류를 흘릴 때 이온이 이동하는 극으로, 전하 값에 따라 자동으로 계산됩니다."),
             ("학습 범위", "'핵심'은 학습 사이트의 기본 범위이자 4쪽 학습지에 실린 항목, '확장'은 수업에서 자주 쓰여 덧붙인 항목입니다."),
             ("교과서 확인", "Ⅳ. 물질의 구성 단원(134~165쪽)에서 실제로 나오는 쪽수를 적었습니다. '–'는 이 단원에 나오지 않는 항목입니다."),
             ("색 구분", "연노란색 행 = 학습 핵심 · 연보라색 행 = 다원자 이온 · 파란 글자 = 양이온, 빨간 글자 = 음이온 · 초록 = 원소·개수비 1 : 1, 주황 = 개수비 다름")]:
    cover.merge_cells(start_row=r, start_column=3, end_row=r, end_column=5)
    write_row(cover, r, [k, v, None, None], height=22, fonts={0: font(10, True, NAVY), 1: font(10)}, aligns={1: LEFT})
    r += 1
notes(cover, r + 1, ["※ 자료 출처: 2022 개정 중학교 과학 교과서 「Ⅳ. 물질의 구성」(134~165쪽) 본문 · 그림 · 표 전체.",
                     "※ 학습 사이트: 송쌤과학 '꼭 알아둘 화학 기호와 화학식' — 이 파일과 같은 자료로 문제를 냅니다.",
                     "※ '이온의 전자 수', '이동하는 극', '전하 합'은 수식이므로 값을 직접 지우지 마세요."], 5)

wb.active = 0
os.makedirs(OUT_DIR, exist_ok=True)
wb.save(OUT)
if __name__ == "__main__":
    print(f"XLSX: {len(wb.sheetnames)}개 시트, {os.path.getsize(OUT)//1024} KB → {OUT}")
